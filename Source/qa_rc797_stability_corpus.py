from pathlib import Path
import json
from openpyxl import load_workbook

root = Path(__file__).resolve().parent
corpus = json.loads((root / "training_corpus.json").read_text(encoding="utf-8"))
main = (root / "main.go").read_text(encoding="utf-8")
seed = (root / "demo_seed.go").read_text(encoding="utf-8")
alternatives = (root / "alternative_scenarios.go").read_text(encoding="utf-8")
reports = (root / "teaching_alternative_scenarios.go").read_text(encoding="utf-8")
transition = main[main.index("func activatePage("):main.index("\nfunc switchPage(")]
out = root.parent / "Training"
lab = load_workbook(out / "03_Лаборатория_20_скважин.xlsx", read_only=True, data_only=False)
economy = load_workbook(out / "04_Экономика_20_скважин.xlsx", read_only=True, data_only=False)

checks = {
    "exactly 20 synthetic wells": corpus["synthetic"] and len(corpus["wells"]) == 20,
    "three failures per well": all(len(w["failures"]) == 3 for w in corpus["wells"]),
    "nine laboratory passes per well": all(len(w["labPasses"]) == 9 for w in corpus["wells"]),
    "every lab pass references a real failure ID": all(p["failureId"] in {e["id"] for e in w["failures"]} for w in corpus["wells"] for p in w["labPasses"]),
    "all three years have repeated lab monitoring": all(all(sum(p["date"].startswith(y) for p in w["labPasses"]) >= 2 for y in ("2024", "2025", "2026")) for w in corpus["wells"]),
    "lab workbook has 180 indexed and full passports": lab["Паспорта"].max_row == 184 and lab["Полные паспорта"].max_row == 184,
    "economy workbook has 180 scenario rows and formulas": economy["Сценарии"].max_row == 184 and str(economy["Сценарии"]["P5"].value).startswith("="),
    "page transitions retain old frame through refresh and target layout": "hidePageControls(previousPage)" in transition and transition.index("refreshPage(id)") < transition.index("layoutCurrentPage()") < transition.index("hidePageControls(previousPage)") < transition.index("commitPageVisibility(id)"),
    "tab switches skip full-window layout after initial layout": "if !mainLayoutReady" in transition and "mainLayoutReady = true" in (root / "stability_frame_windows.go").read_text(encoding="utf-8"),
    "page navigation does not intercept minimize or scan all children": "SC_MINIMIZE" not in main and "sweepInactivePageControls" not in main and "pEnumChildWindows.Call" not in main,
    "no workflow transaction toggles redraw on the top-level window": all("pSendMessageW.Call(uintptr(hMain), WM_SETREDRAW" not in (root / name).read_text(encoding="utf-8") for name in ("laboratory.go", "methods_workflow.go")),
    "page show calls defer physical visibility during transitions": "if pageSwitching" in main[main.index("func show(h syscall.Handle, yes bool)"):main.index("\nfunc beginBusyLoading(")],
    "lab seeding uses explicit failure links": "failuresByID[p.FailureID]" in seed,
    "obsolete pass-count guard removed": "if i >= len(w.Failures)" not in seed,
    "obsolete sample decisions are excluded from active history and object summaries": "isRetiredLegacyTeachingObjectKey(record.ObjectKey)" in main and "isRetiredLegacyTeachingObjectKey(v.ObjectKey)" in main,
    "old alternatives are retired when shared records load": "isRetiredLegacyTeachingObjectKey(scenario.ObjectKey)" in alternatives,
    "final report archive creates current records for all 20 wells": "for ordinal := 1; ordinal <= 20; ordinal++" in reports and 'ID: id, CreatedAt: created, ObjectKey: objectKey' in reports,
    "canonical lab IDs are validated and matching covers the full passport": "trainingLabPassMatchesCanonical(old, expected)" in seed and "canonicalTrainingPassIDs()" in (root / "training_corpus_app.go").read_text(encoding="utf-8"),
    "startup, off-page and preparing controls are denied at the native visibility boundary": "pageControlMayBeVisible(owner, currentPage, yes, preparing)" in main and "if activePage != currentPage" in main,
    "startup reconciles all registered controls before dashboard activation": main.index("buildLaboratoryUI()") < main.index("for page := range pageControls") < main.index("activatePage(ID_DASH, false)"),
    "release source version is 807-RC": 'const appBuildVersion = "7.3.11.807-RC"' in main,
}
for label, passed in checks.items():
    print(("PASS" if passed else "FAIL") + ": " + label)
if not all(checks.values()):
    raise SystemExit(1)
