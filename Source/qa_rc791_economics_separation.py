from pathlib import Path
from openpyxl import load_workbook

root = Path(__file__).resolve().parent
out = root / 'outputs' / 'rosa_training_20_20261008'
failures = load_workbook(out / '02_Реестр_отказов_20_скважин.xlsx', read_only=True, data_only=False)
economy = load_workbook(out / '04_Экономика_20_скважин.xlsx', read_only=True, data_only=False)
template = load_workbook(root / 'РОСА_БЛАНК_ЭКОНОМИКИ.xlsx', read_only=True, data_only=False)
main = (root / 'main.go').read_text(encoding='utf-8')
flow = (root / 'workflow_pages.go').read_text(encoding='utf-8')
econ_ui = (root / 'economics_visual_reference.go').read_text(encoding='utf-8')
source = (root / 'economics_template_windows.go').read_text(encoding='utf-8')
notice = (root / 'section_banner_settings.go').read_text(encoding='utf-8')

checks = {
    'failure register has no economics sheet': all('эконом' not in name.lower() for name in failures.sheetnames),
    'separate economy workbook has its own summary': 'Сводка экономики' in economy.sheetnames and 'История отказов' not in economy.sheetnames,
    'economy workbook contains 180 scenario records': economy['Сценарии'].max_row == 184,
    'economy scenario rows carry spreadsheet formulas': str(economy['Сценарии']['J5'].value).startswith('=') and str(economy['Сценарии']['P5'].value).startswith('='),
    'economics template has separate input, scenarios, and comparison sheets': template.sheetnames == ['Параметры', 'Варианты', 'Сравнение'],
    'economics comparison uses formulas': str(template['Сравнение']['H5'].value).startswith('='),
    'section 6 exposes download and upload controls': 'ID_ECON_TEMPLATE_SAVE' in flow and 'ID_ECON_TEMPLATE_LOAD' in flow and 'Скачать бланк' in econ_ui and 'Загрузить бланк' in econ_ui,
    'upload validates the selected object and changes only economics scenarios': 'laboratoryObjectKeysEquivalent(imported.objectKey, selectedObjectKey)' in source and 'saveAlternativeScenarios()' in source and 'Реестр отказов не затронут' in source,
    'attention cue is text only and blinks': 'sectionNoticeIcon' not in main + notice and 'sectionNoticePulse' in main and 'textColor = rgb(255, 237, 240)' in main,
    'single click opens the teaching-data warning': 'idSectionNoticeLabel && code == 0' in main and 'sectionNoticeMessage' in notice,
    'release version advanced': 'const appBuildVersion = "7.3.11.791-RC"' in main,
}
for label, passed in checks.items():
    print(('PASS' if passed else 'FAIL') + ': ' + label)
if not all(checks.values()):
    raise SystemExit(1)
