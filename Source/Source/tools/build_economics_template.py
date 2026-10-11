from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'РОСА_БЛАНК_ЭКОНОМИКИ.xlsx'
NAVY='163A5B'; RED='B42318'; TEAL='0F6B78'; PALE='EEF3F8'; RED_PALE='FCE8E8'; WHITE='FFFFFF'; BORDER='C9D4E0'; INPUT='FFF4D6'
thin=Side(style='thin', color=BORDER)
wb=Workbook()
ws=wb.active; ws.title='Параметры'
ws.append(['РОСА · БЛАНК ЭКОНОМИЧЕСКОЙ ОЦЕНКИ'])
ws.append(['Заполните выделенные ячейки. Загружайте файл в разделе «6. Экономическая оценка» для выбранного объекта.'])
ws.append(['Учебные значения и формулы не заменяют исходные данные и проверку экономиста. Реестр отказов в этот бланк не включается.'])
ws.append(['Параметр','Значение','Единица','Пояснение'])
parameters=[
 ('Объект','', '', 'Должен совпадать с объектом, закреплённым в РОСА; пустое значение означает текущий объект.'),
 ('Текущий МРП',None,'сут.','Фактическая или принятая база для сравнения.'),
 ('МРП осторожный',None,'сут.','Нижний сценарный прогноз.'),
 ('МРП центральный',None,'сут.','Основной прогноз на выбранный горизонт.'),
 ('МРП благоприятный',None,'сут.','Верхний сценарный прогноз.'),
 ('Выбранный сценарий','Центральный','','Осторожный, Центральный или Благоприятный.'),
 ('Горизонт',36,'мес.','Целое число от 1 до 60 месяцев.'),
 ('Длительность ремонта',None,'сут.','Средняя длительность одного ремонта.'),
 ('Стоимость ремонта',None,'руб.','Стоимость одного ремонта.'),
 ('OPEX',None,'руб./год','Эксплуатационные затраты выбранного варианта.'),
 ('Инфляция OPEX',0,'%/год','Годовой темп изменения эксплуатационных затрат.'),
 ('CAPEX',None,'руб.','Разовые капитальные и материальные затраты.')]
for p in parameters: ws.append(list(p))
for row in ws.iter_rows(min_row=1,max_row=ws.max_row,max_col=4):
 for c in row:
  c.alignment=Alignment(vertical='center',wrap_text=True)
  if c.row>=4: c.border=Border(bottom=thin)
for c in ws[1]: c.fill=PatternFill('solid',fgColor=NAVY);c.font=Font(name='Segoe UI',size=15,bold=True,color=WHITE)
ws.merge_cells('A1:D1')
for rr,color in ((2,PALE),(3,RED_PALE)):
 for c in ws[rr]: c.fill=PatternFill('solid',fgColor=color);c.font=Font(name='Segoe UI',size=10,bold=(rr==3),color=(RED if rr==3 else NAVY))
ws.merge_cells('A2:D2');ws.merge_cells('A3:D3')
for c in ws[4]: c.fill=PatternFill('solid',fgColor=TEAL);c.font=Font(name='Segoe UI',bold=True,color=WHITE)
for r in range(5,ws.max_row+1):
 ws.cell(r,2).fill=PatternFill('solid',fgColor=INPUT)
 for c in ws[r]: c.font=Font(name='Segoe UI',size=10,color=NAVY)
ws.column_dimensions['A'].width=30;ws.column_dimensions['B'].width=22;ws.column_dimensions['C'].width=17;ws.column_dimensions['D'].width=72
ws.row_dimensions[1].height=28;ws.row_dimensions[2].height=30;ws.row_dimensions[3].height=34;ws.row_dimensions[4].height=25
for r in range(5,ws.max_row+1): ws.row_dimensions[r].height=30
ws.freeze_panes='B5';ws.auto_filter.ref=f'A4:D{ws.max_row}';ws.sheet_view.showGridLines=False
scenario=DataValidation(type='list',formula1='"Осторожный,Центральный,Благоприятный"',allow_blank=False)
ws.add_data_validation(scenario);scenario.add(ws['B10'])
integer=DataValidation(type='whole',operator='between',formula1=1,formula2=60,allow_blank=False)
ws.add_data_validation(integer);integer.add(ws['B11'])

variants=wb.create_sheet('Варианты')
variants.append(['ПАРАМЕТРЫ ПАРАЛЛЕЛЬНЫХ ЭКОНОМИЧЕСКИХ ВАРИАНТОВ'])
variants.append(['Не меняйте состав методов в этой таблице. Названия и методы берутся из раздела 4; при загрузке строки применяются к вариантам по порядку.'])
variants.append(['Пустые значения не заменяют показатели, уже сохранённые в выбранной схеме. Суммы — в рублях, МРП — в сутках.'])
variants.append(['№','Вариант','Методы (справочно)','МРП осторожный','МРП центральный','МРП благоприятный','OPEX, руб./год','CAPEX, руб.'])
for i in range(1,4): variants.append([i,f'Вариант {i}','',None,None,None,None,None])
for row in variants.iter_rows(min_row=1,max_row=variants.max_row,max_col=8):
 for c in row: c.alignment=Alignment(vertical='center',wrap_text=True)
for c in variants[1]:c.fill=PatternFill('solid',fgColor=NAVY);c.font=Font(name='Segoe UI',size=14,bold=True,color=WHITE)
variants.merge_cells('A1:H1')
for rr,color in ((2,PALE),(3,RED_PALE)):
 for c in variants[rr]:c.fill=PatternFill('solid',fgColor=color);c.font=Font(name='Segoe UI',size=10,bold=(rr==3),color=(RED if rr==3 else NAVY))
variants.merge_cells('A2:H2');variants.merge_cells('A3:H3')
for c in variants[4]: c.fill=PatternFill('solid',fgColor=TEAL);c.font=Font(name='Segoe UI',bold=True,color=WHITE)
for r in range(5,8):
 for c in variants[r]:c.border=Border(bottom=thin);c.font=Font(name='Segoe UI',size=10,color=NAVY)
 for col in range(4,9):variants.cell(r,col).fill=PatternFill('solid',fgColor=INPUT)
for col,width in enumerate([9,22,48,20,20,22,22,20],1):variants.column_dimensions[get_column_letter(col)].width=width
variants.row_dimensions[1].height=28;variants.row_dimensions[2].height=30;variants.row_dimensions[3].height=30;variants.row_dimensions[4].height=36
for r in range(5,8):variants.row_dimensions[r].height=28
variants.freeze_panes='D5';variants.auto_filter.ref='A4:H7';variants.sheet_view.showGridLines=False

calc=wb.create_sheet('Сравнение')
calc.append(['СРАВНЕНИЕ ВАРИАНТОВ · ФОРМУЛЫ'])
calc.append(['Результаты пересчитываются в Excel после заполнения листов «Параметры» и «Варианты». Суммарные значения не включают реестр отказов.'])
calc.append([])
calc.append(['Вариант','МРП для расчёта, сут.','Предотвращено ремонтов, шт./год','Предотвращённый простой, сут./год','Валовая выгода за горизонт, руб.','OPEX за горизонт, руб.','CAPEX, руб.','Чистая выгода, руб.','Выгода/затраты','Окупаемость, мес.'])
for i,vr in enumerate(range(5,8),5):
 calc.append([f'=Варианты!B{vr}',f'=IF(Варианты!E{vr}="",Параметры!$B$8,Варианты!E{vr})',f'=IFERROR(MAX(0,365/Параметры!$B$6-365/B{i}),0)',f'=C{i}*IF(Параметры!$B$12="",0,Параметры!$B$12)',f'=C{i}*Параметры!$B$13*Параметры!$B$11/12',f'=SUMPRODUCT((ROW($1:$60)<=Параметры!$B$11)*IF(Варианты!G{vr}="",Параметры!$B$14,Варианты!G{vr})/12*(1+Параметры!$B$15/100)^INT((ROW($1:$60)-1)/12))',f'=IF(Варианты!H{vr}="",Параметры!$B$16,Варианты!H{vr})',f'=E{i}-F{i}-G{i}',f'=IFERROR(E{i}/(F{i}+G{i}),0)',f'=IFERROR(G{i}/(C{i}*Параметры!$B$13-F{i}*12/Параметры!$B$11)*12,"")'])
for row in calc.iter_rows(min_row=1,max_row=calc.max_row,max_col=10):
 for c in row:c.alignment=Alignment(vertical='center',wrap_text=True);c.border=Border(bottom=thin)
for c in calc[1]:c.fill=PatternFill('solid',fgColor=NAVY);c.font=Font(name='Segoe UI',size=14,bold=True,color=WHITE)
calc.merge_cells('A1:J1')
for c in calc[2]:c.fill=PatternFill('solid',fgColor=PALE);c.font=Font(name='Segoe UI',size=10,color=NAVY)
calc.merge_cells('A2:J2')
for c in calc[4]:c.fill=PatternFill('solid',fgColor=TEAL);c.font=Font(name='Segoe UI',size=10,bold=True,color=WHITE)
for r in range(5,8):
 for c in calc[r]:c.font=Font(name='Segoe UI',size=10,color=NAVY)
for col,width in enumerate([25,20,23,24,26,22,18,23,17,19],1):calc.column_dimensions[get_column_letter(col)].width=width
calc.row_dimensions[1].height=28;calc.row_dimensions[2].height=38;calc.row_dimensions[4].height=48
for r in range(5,8):calc.row_dimensions[r].height=34
calc.freeze_panes='A5';calc.sheet_view.showGridLines=False
calc.sheet_properties.pageSetUpPr.fitToPage=True;calc.page_setup.fitToWidth=1;calc.page_setup.fitToHeight=1;calc.page_setup.orientation='landscape'

wb.calculation.fullCalcOnLoad=True;wb.calculation.forceFullCalc=True;wb.calculation.calcMode='auto'
wb.save(OUT)
print(OUT)
