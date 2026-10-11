import json, re, math, zipfile
from pathlib import Path
from datetime import date, timedelta
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs'/'rosa_training_20_20261009'; OUT.mkdir(parents=True,exist_ok=True)
fields=[('Авиловское учебное','Бобриковский горизонт'),('Бешкульское учебное','Батский III пласт'),('Северное учебное','Песчаный горизонт'),('Южное учебное','Турнейский горизонт'),('Заречное учебное','Визейский горизонт'),('Речное учебное','Девонский горизонт'),('Степное учебное','Карбонатный пласт'),('Лесное учебное','Тульский горизонт'),('Западное учебное','Бобриковский горизонт'),('Восточное учебное','Песчаный горизонт'),('Приозёрное учебное','Батский III пласт'),('Ключевое учебное','Турнейский горизонт'),('Майское учебное','Визейский горизонт'),('Октябрьское учебное','Девонский горизонт'),('Сосновское учебное','Карбонатный пласт'),('Луговое учебное','Тульский горизонт'),('Берёзовское учебное','Бобриковский горизонт'),('Уральское учебное','Песчаный горизонт'),('Верхнее учебное','Батский III пласт'),('Юго-Западное учебное','Турнейский горизонт')]
lifts=['УЭЦН','УШГН','Фонтанный','УЭВН','ОРД']*4
compl=[('Солеотложения','С2','Карбонатный осадок; рост перепада давления','Насос / НКТ','scale','Ингибитор солеотложений'),('Коррозия','К2','Локальное истончение металла; повышенное железо','НКТ / муфтовое соединение','corrosion','Ингибитор коррозии + контроль остатка'),('АСПО','А2','Парафиновые отложения; падение подачи','Приёмная часть насоса','aspo','Подбор растворителя и профилактики'),('Механические примеси','М2','Рост КВЧ и абразивный износ','Рабочая пара / фильтр','solids','Фильтрация и контроль фракций'),('Свободный газ','Г1','Нестабильная подача; повышенный газовый фактор','Газосепаратор','gases','Подбор газосепаратора'),('Высокая вязкость / эмульсия','В2','Рост вязкости и замедленное водоотделение','Насос / выкидная линия','aspo','Подбор деэмульгатора'),('Биокоррозия','К1','Микробиологический риск и рост железа','НКТ / защитное покрытие','microbiology','Биоцидная обработка + контроль'),('Гидраты','Гд1','Низкая температура и влажный газ','Устьевая арматура / линия','gases','Тепловой режим + ингибитор гидратов')]
def lab_value(key,n,j):
 seed=n*23+j*17+len(key)*11
 special={'ph':lambda:round(6.1+((n+j)%9)*.13,2),'chlorides':lambda:9000+(n*911+j*1301)%33000,'cl':lambda:9000+(n*911+j*1301)%33000,'ca':lambda:280+(n*37+j*61)%820,'mg':lambda:60+(n*19+j*31)%260,'na_k':lambda:9000+(n*317+j*541)%19000,'so4':lambda:240+(n*23+j*47)%1300,'hco3':lambda:180+(n*29+j*43)%980,'co2':lambda:18+(n*29+j*41)%410,'h2s':lambda:round(((n*11+j*17)%420)/100,2),'sulfide':lambda:round(((n*11+j*17)%420)/100,2),'o2':lambda:round(((n*3+j*7)%28)/100,2),'water':lambda:25+(n*7+j*13)%65,'paraffins':lambda:round(3+((n*7+j*13)%105)/10,1),'resins':lambda:round(5+((n*9+j*11)%115)/10,1),'asphaltenes':lambda:round(1+((n*5+j*7)%49)/10,1),'viscosity':lambda:32+(n*17+j*31)%270,'mineralization':lambda:round(28+((n*13+j*17)%420)/10,1),'iron':lambda:round(1.2+((n*11+j*17)%140)/10,1),'iron_total':lambda:round(1.2+((n*11+j*17)%140)/10,1),'kvch':lambda:45+(n*17+j*29)%460,'suspended':lambda:22+(n*19+j*37)%780,'sand':lambda:round(((n*7+j*11)%30)/100,2),'gas_factor':lambda:28+(n*31+j*47)%330,'methane':lambda:round(72+((n*3+j*7)%270)/10,1)}
 if key in special:return special[key]()
 return f'Учебное синтетическое значение {seed%997}'
wells=[]
for i,(field,layer) in enumerate(fields):
 n=i+1; liquid=42+(i*17)%71; wc=23+(i*11+19)%64; oil=round(liquid*(100-wc)/100,1); name=f'Учебная скважина {n:02d} · {field} / скв. {101+i} / {layer}'
 w=dict(id=f'TRN-{n:02d}',ordinal=n,objectKey=name,field=field,well=str(101+i),layer=layer,liftMethod=lifts[i],gno=lifts[i],pumpModel=f'{lifts[i]} · учебный типоразмер',pumpDepth=1080+(i*37%520),liquidRate=liquid,waterCut=wc,oilRate=oil,netbackRubPerM3=[11500,12800,10400,12100,9800][i%5],reservoirPressure=round(12.5+(i%7)*.7,1),bottomholePressure=round(7.2+(i%6)*.55,1),failures=[],labPasses=[])
 for j in range(3):
  c=compl[(i+j*3+(i%3))%len(compl)]; year=2024+j; month=2+((i*3+j*5)%10); day=1+((i*7+j*11)%26); dt=date(year,month,min(day,28)); dt=min(dt,date(2026,10,8)); start=date(2023,11,1+i%10) if j==0 else date.fromisoformat(w['failures'][-1]['repairEnd']); mrp=max(45,(dt-start).days); repair=dt+timedelta(days=4+(i+j)%7); repeat=(i+j)%4==0 or (j==2 and i%3==1); method=c[5] if not repeat else 'Повторная корректировка: '+c[5]
  e=dict(id=f'{w["id"]}-F{j+1}',eventNo=j+1,date=dt.isoformat(),dateRu=dt.strftime('%d.%m.%Y'),cycleStart=start.isoformat(),cycleStartRu=start.strftime('%d.%m.%Y'),repairEnd=repair.isoformat(),repairEndRu=repair.strftime('%d.%m.%Y'),mrp=mrp,liftMethod=lifts[i],equipment=lifts[i],brand=f'Учебное исполнение {(i%6)+1:02d}',size=w['pumpModel'],pumpDepth=w['pumpDepth'],complication=c[0],classification=c[1],associatedComplication=compl[(i+4)%8][0] if j==1 else '',cause=c[2],failedUnit=c[3],methodBefore='Штатный режим / предыдущее мероприятие',methodAfter=method,decisionType='Повтор с корректировкой' if repeat else 'Новый подбор',labGroup=c[4],labPassId=f'{w["id"]}-L{j+1}',investigation='Учебный разбор; лабораторные признаки и дефектовка смоделированы.',commission=f'Учебное заключение по событию {j+1}; не производственный акт.',resultingMRP=mrp+25+((i*13+j*17)%55),comment='Синтетическая учебная запись.')
  w['failures'].append(e); w['labPasses'].append(dict(passId=e['labPassId'],date=(dt-timedelta(days=3)).isoformat(),dateRu=(dt-timedelta(days=3)).strftime('%d.%m.%Y'),failureId=e['id'],complication=c[0],groupFocus=c[4],conclusion=f'Синтетический паспорт перед событием {j+1}: {c[0]}',status='saved'))
 # Two planned monitoring passes per year supplement the pre-failure passports.
 # Each pass is linked to the next recorded failure (or the latest event at the
 # end of the period), so lab trends, failure history and economics share IDs.
 for year in (2024,2025,2026):
  for month,day in ((5,15),(9,15)):
   sample=date(year,month,day)
   linked=min((e for e in w['failures'] if date.fromisoformat(e['date'])>=sample),key=lambda e:e['date'],default=w['failures'][-1])
   j=linked['eventNo']
   c=compl[(i+(j-1)*3+(i%3))%len(compl)]
   w['labPasses'].append(dict(passId=f'{w["id"]}-LM-{year}-{month:02d}',date=sample.isoformat(),dateRu=sample.strftime('%d.%m.%Y'),failureId=linked['id'],complication=linked['complication'],groupFocus=linked['labGroup'],conclusion=f'Плановый синтетический мониторинг {sample.strftime("%m.%Y")}; связь с событием {linked["eventNo"]}: {linked["complication"]}',status='saved',monitoring=True))
 w['labPasses'].sort(key=lambda p:(p['date'],p['passId']))
 wells.append(w)
corpus=dict(schemaVersion=1,title='Учебный фонд РОСА — 20 скважин',synthetic=True,periodStart='2023-11-01',periodEnd='2026-10-08',wells=wells)
(ROOT/'training_corpus.json').write_text(json.dumps(corpus,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

NAVY='163A5B'; TEAL='0F6B78'; RED='B42318'; THIN=Side(style='thin',color='D5DEE8')
def style(ws,freeze='A5'):
 ws.freeze_panes=freeze;ws.sheet_view.showGridLines=False;ws.auto_filter.ref=ws.dimensions
 for c in ws[1]: c.font=Font(name='Arial',size=14,bold=True,color=NAVY)
 for row in ws.iter_rows(min_row=4):
  for c in row:
   c.alignment=Alignment(vertical='top',wrap_text=True);c.border=Border(bottom=THIN)
   if c.row==4:c.fill=PatternFill('solid',fgColor=NAVY);c.font=Font(name='Arial',bold=True,color='FFFFFF')

# Tech regime and three-year event timeline
wb=Workbook();s=wb.active;s.title='Техрежим';h=wb.create_sheet('История 36 мес.')
s.append(['Учебный техрежим · 20 скважин · снимок на 08.10.2026']);s.append(['Синтетические данные; использовать только для обучения.'])
s.append([]);s.append(['ID','Объект','Месторождение','Скважина','Пласт','Способ добычи','ГНО','Глубина насоса, м','Дебит жидкости, м³/сут','Обводнённость, %','Дебит нефти, м³/сут'])
for w in wells:s.append([w['id'],w['objectKey'],w['field'],w['well'],w['layer'],w['liftMethod'],w['gno'],w['pumpDepth'],w['liquidRate'],w['waterCut'],w['oilRate']])
h.append(['История техрежима и отказов · 01.11.2023–08.10.2026']);h.append(['Три хронологические точки на каждую скважину; ID связывает отказ, анализ и экономику.']);h.append([]);h.append(['ID','Объект','Дата отказа','ID отказа','МРП, сут.','Способ добычи','ГНО','Осложнение','Метод до','Метод после','Дебит нефти','Обводнённость','Статус'])
for w in wells:
 for e in w['failures']:h.append([w['id'],w['objectKey'],e['dateRu'],e['id'],e['mrp'],w['liftMethod'],w['gno'],e['complication'],e['methodBefore'],e['methodAfter'],w['oilRate'],w['waterCut'],'Учебные синтетические данные'])
for ws in (s,h):style(ws);ws.row_dimensions[1].height=26;ws.row_dimensions[2].height=30
for col,width in {'A':16,'B':50,'C':24,'D':14,'E':25,'F':19,'G':15,'H':20,'I':20,'J':18,'K':20,'L':18,'M':30}.items():s.column_dimensions[col].width=width;h.column_dimensions[col].width=width
wb.save(OUT/'01_Техрежим_20_скважин.xlsx')

# Failure template with populated rows; economy is formula-based with alternatives.
wb=load_workbook(ROOT/'РОСА_БЛАНК_РЕЕСТРА_ОТКАЗОВ_632-RC.xlsx'); sh=wb['История отказов']; op=wb['Оперативный ввод']; headers=[sh.cell(4,c).value for c in range(1,71)]; opheaders=[op.cell(4,c).value for c in range(1,32)]
labels={str(x).replace('[ДОП]','').strip().lower():i+1 for i,x in enumerate(headers) if x}
for idx,(w,e) in enumerate([(w,e) for w in wells for e in w['failures']],5):
  vals={'Месторождение*':w['field'],'№ скважины*':w['well'],'Дата пуска*':e['cycleStartRu'],'Дата отказа / окончания*':e['dateRu'],'МРП, сут. (авто)':e['mrp'],'Способ эксплуатации*':w['liftMethod'],'Типоразмер ГНО 1*':w['pumpModel'],'Марка ГНО*':e['brand'],'Тип события*':'Отказ','Признак / предпосылка*':e['cause'],'Отказавший узел*':e['failedUnit'],'Предварительная причина*':e['cause'],'Основная категория осложнения*':e['complication'],'Сопутствующее осложнение 1':e['associatedComplication'],'ТКРС 1':f'Учебное мероприятие №{e["eventNo"]}','Вид повреждения':e['failedUnit'],'Итоговая причина':e['cause'],'Уточнённая категория осложнения':e['classification'],'Метод до отказа 1':e['methodBefore'],'Метод после отказа 1':e['methodAfter'],'Подтверждение':f'Лабораторный паспорт {e["labPassId"]}','Статус расследования':'Учебный разбор заполнен','Достоверность':'Синтетический учебный кейс','Заключение комиссии / акт':e['commission'],'Источник / документ':f'Учебный паспорт {e["labPassId"]}','Тип насоса / способ эксплуатации':w['pumpModel'],'Глубина насоса, м':w['pumpDepth'],'Условный диаметр / размер ШГН, мм':'Учебное значение','Производитель насоса':'Учебный изготовитель','Тип НКТ':'Учебный НКТ','Длина НКТ, м':w['pumpDepth']+120,'Год НКТ':2020+e['eventNo'],'СПО НКТ':'Учебная спуско-подъёмная операция','Состояние двигателя до спуска':'Учебное состояние','Производитель двигателя':'Учебный изготовитель','Типоразмер двигателя (габарит), мм':'Учебное исполнение','Год двигателя':2020+e['eventNo'],'Состояние кабеля до спуска':'Учебное состояние','Сечение кабеля, мм²':25,'Длина кабеля, м':w['pumpDepth']+30,'Состояние удлинителя до спуска':'Учебное состояние','Сечение удлинителя, мм²':16,'Длина удлинителя, м':w['pumpDepth']+15,'Состояние гидрозащиты до спуска':'Учебное состояние','Производитель гидрозащиты':'Учебный изготовитель','Типоразмер гидрозащиты (габарит), мм':'Учебное исполнение','Год гидрозащиты':2020+e['eventNo'],'Состояние ТМС до спуска':'Учебное состояние','Производитель ТМС':'Учебный изготовитель','Типоразмер ТМС':'Учебное исполнение','Год ТМС':2020+e['eventNo'],'Состояние штанг до спуска':'Учебное состояние','Производитель штанг':'Учебный изготовитель','Штанги Ø19 мм, м':int((w['pumpDepth']+120)*.2),'Штанги Ø22 мм, м':int((w['pumpDepth']+120)*.4),'Штанги Ø25 мм, м':int((w['pumpDepth']+120)*.4),'Год выпуска штанг':2020+e['eventNo'],'СПО штанг':'Учебная операция','Комментарий':f"{e['id']}; {e['decisionType']}; {e['comment']}"}
  for k,v in vals.items():
   c=labels.get(k.lower())
   if c:sh.cell(idx,c).value=v
  for c,hdr in enumerate(opheaders,1):
   if not hdr: continue
   key=str(hdr).replace('[ДОП]','').strip().lower(); value=vals.get(key)
   if value is None:
    value={'статус предварительного (авто)':'Предварительное заполнено синтетически','предварительное заключение (авто)':f"{e['complication']} — учебная предварительная оценка"}.get(key,'Учебное значение заполнено для проверки шаблона')
   if op.cell(idx,c).data_type!='f': op.cell(idx,c).value=value
  for c,hdr in enumerate(headers,1):
   if not hdr: continue
   cell=sh.cell(idx,c)
   if cell.data_type=='f': continue
   key=str(hdr).replace('[ДОП]','').strip().lower()
   if cell.value is None: cell.value=vals.get(key,'Учебное значение заполнено для проверки полного шаблона')
# Keep the failure registry and economic scenarios in separate workbooks.
# The refusal file contains only failure records; economic calculations have
# their own workbook, summaries and scenario-level effect statistics.
style(sh)
wb.save(OUT/'02_Реестр_отказов_20_скважин.xlsx')

econ_wb=Workbook();econ=econ_wb.active;econ.title='Сценарии'
econ.append(['Экономические варианты · 20 скважин · три цикла по каждому объекту'])
econ.append(['Сценарии и суммы синтетические; агрегаты ниже не являются фактическим производственным эффектом.'])
econ.append([])
econ.append(['ID экономического варианта','ID скважины','Объект','Дата оценки','Цикл, №','Сценарий','Вариант мероприятия','МРП до, сут.','Прогноз МРП, сут.','Доп. дни','Дебит нефти, м³/сут','Netback, руб./м³','Валовый эффект, руб.','CAPEX, руб.','OPEX, руб./год','Чистый эффект, руб.','Выгода/затраты','Окупаемость, мес.'])
for w in wells:
 for e in w['failures']:
  for k,scenario in enumerate(['Повтор метода','Новый подбор','Комплексная защита']):
   r=econ.max_row+1;target=e['resultingMRP']+[0,30,60][k];capex=[280000,520000,760000][k]+w['ordinal']*3500;opex=[115000,165000,235000][k]
   assessment_date=min(date.fromisoformat(e['repairEnd']),date(2026,10,8))
   variant=[e['methodAfter'],e['methodAfter']+' · лабораторная корректировка','Комплекс: '+e['methodAfter']+' + мониторинг'][k]
   econ_id=f"EC-{w['id']}-C{e['eventNo']}-{chr(65+k)}"
   econ.append([econ_id,w['id'],w['objectKey'],assessment_date,e['eventNo'],scenario,variant,e['mrp'],target,f'=MAX(I{r}-H{r},0)',w['oilRate'],w['netbackRubPerM3'],f'=J{r}*K{r}*L{r}',capex,opex,f'=M{r}-N{r}-O{r}*J{r}/365',f'=IFERROR(M{r}/(N{r}+O{r}*J{r}/365),0)',f'=IFERROR(N{r}/((M{r}/MAX(1,J{r})*365)-O{r})*12,0)'])
summary=econ_wb.create_sheet('Сводка экономики')
summary.append(['Экономическая сводка по сценариям'])
summary.append(['Каждая строка сравнивает условный вариант для одного цикла. Суммарные значения — учебная выборка, не прогноз бюджета фонда.'])
summary.append([])
summary.append(['Сценарий','Вариантов','Доп. дни МРП','Валовый эффект, руб.','CAPEX, руб.','OPEX с учетом доп. дней, руб.','Чистый эффект, руб.','Средний эффект на вариант, руб.'])
last=econ.max_row
scenarios=['Повтор метода','Новый подбор','Комплексная защита']
for i,scenario in enumerate(scenarios,5):
 summary.append([scenario,f'=COUNTIF(Сценарии!$F$5:$F${last},A{i})',f'=SUMIF(Сценарии!$F$5:$F${last},A{i},Сценарии!$J$5:$J${last})',f'=SUMIF(Сценарии!$F$5:$F${last},A{i},Сценарии!$M$5:$M${last})',f'=SUMIF(Сценарии!$F$5:$F${last},A{i},Сценарии!$N$5:$N${last})',f'=SUMPRODUCT((Сценарии!$F$5:$F${last}=A{i})*Сценарии!$O$5:$O${last}*Сценарии!$J$5:$J${last}/365)',f'=D{i}-E{i}-F{i}',f'=IFERROR(G{i}/B{i},0)'])
summary.append([]);summary.append(['Динамика экономических сценариев по году оценки']);summary.append([])
summary.append(['Год','Сценарий','Вариантов','Доп. дни МРП','Валовый эффект, руб.','CAPEX, руб.','OPEX, руб.','Чистый эффект, руб.'])
for year in (2024,2025,2026):
 for scenario in scenarios:
  r=summary.max_row+1;summary.append([year,scenario,f'=COUNTIFS(Сценарии!$F$5:$F${last},B{r},Сценарии!$D$5:$D${last},">="&DATE(A{r},1,1),Сценарии!$D$5:$D${last},"<="&DATE(A{r},12,31))',f'=SUMIFS(Сценарии!$J$5:$J${last},Сценарии!$F$5:$F${last},B{r},Сценарии!$D$5:$D${last},">="&DATE(A{r},1,1),Сценарии!$D$5:$D${last},"<="&DATE(A{r},12,31))',f'=SUMIFS(Сценарии!$M$5:$M${last},Сценарии!$F$5:$F${last},B{r},Сценарии!$D$5:$D${last},">="&DATE(A{r},1,1),Сценарии!$D$5:$D${last},"<="&DATE(A{r},12,31))',f'=SUMIFS(Сценарии!$N$5:$N${last},Сценарии!$F$5:$F${last},B{r},Сценарии!$D$5:$D${last},">="&DATE(A{r},1,1),Сценарии!$D$5:$D${last},"<="&DATE(A{r},12,31))',f'=SUMPRODUCT((Сценарии!$F$5:$F${last}=B{r})*(Сценарии!$D$5:$D${last}>=DATE(A{r},1,1))*(Сценарии!$D$5:$D${last}<=DATE(A{r},12,31))*Сценарии!$O$5:$O${last}*Сценарии!$J$5:$J${last}/365)',f'=E{r}-F{r}-G{r}'])
for ws in (econ,summary):
 style(ws)
 ws.row_dimensions[1].height=26;ws.row_dimensions[2].height=30
for c,width in enumerate([27,16,52,17,12,25,46,15,18,14,19,18,21,18,18,21,18,20],1):econ.column_dimensions[get_column_letter(c)].width=width
for c,width in enumerate([28,16,18,22,22,28,22,26],1):summary.column_dimensions[get_column_letter(c)].width=width
for c,width in enumerate([12,25,16,18,22,18,18,22],1):summary.column_dimensions[get_column_letter(c)].width=width
for ws in (econ,summary): ws.freeze_panes='A5';ws.sheet_view.zoomScale=85
wb.save(OUT/'02_Реестр_отказов_20_скважин.xlsx')
econ_wb.save(OUT/'04_Экономика_20_скважин.xlsx')

# Lab passports: compact index and full parameter matrix parsed from application catalog.
labsrc=(ROOT/'laboratory.go').read_text(encoding='utf-8'); groups=[]
for gm in re.finditer(r'\{ID:\s*ID_LAB_TAB_[^,]+,\s*Code:\s*"([^"]+)",\s*Name:\s*"([^"]+)",\s*Fields:\s*\[\]LabFieldDef\{([\s\S]*?)\n\s*\}\}',labsrc):
 if gm.group(1)=='overall':continue
 for fm in re.finditer(r'\{Key:\s*"([^"]+)",\s*Label:\s*"([^"]+)",\s*Unit:\s*"([^"]*)"\}',gm.group(3)):groups.append((gm.group(1),gm.group(2),fm.group(1),fm.group(2),fm.group(3)))
lab=Workbook();idx=lab.active;idx.title='Паспорта';full=lab.create_sheet('Полные паспорта')
idx.append(['Индекс лабораторных паспортов · 180 анализов']);idx.append(['Девять анализов на скважину: три перед отказами и шесть плановых контрольных; каждый связан по ID с объектом и событием. Все значения синтетические.']);idx.append([]);idx.append(['ID паспорта','ID отказа','ID скважины','Объект','Дата анализа','Осложнение','Метод после отказа','Вывод','Статус'])
for w in wells:
 by_failure={e['id']:e for e in w['failures']}
 for p in w['labPasses']:
  e=by_failure[p['failureId']]
  idx.append([p['passId'],e['id'],w['id'],w['objectKey'],date.fromisoformat(p['date']).strftime('%d.%m.%Y'),e['complication'],e['methodAfter'],p['conclusion'],'Учебные данные'])
full.append(['Полные лабораторные паспорта · все поля каталога ROSA']);full.append(['Фильтр по ID скважины/паспорта; показатели рассчитаны условно и не являются лабораторными измерениями.']);full.append([]);full.append(['ID паспорта','ID отказа','ID скважины','Объект','Дата анализа','Осложнение','Метод','Статус']+[f'{g} · {label} ({unit})' for code,g,key,label,unit in groups])
for w in wells:
 by_failure={e['id']:e for e in w['failures']}
 for p in w['labPasses']:
  e=by_failure[p['failureId']]
  full.append([p['passId'],e['id'],w['id'],w['objectKey'],date.fromisoformat(p['date']).strftime('%d.%m.%Y'),e['complication'],e['methodAfter'],'Учебное · синтетическое']+[lab_value(key,w['ordinal'],e['eventNo']+(int(p['date'][5:7])//3)) for code,g,key,label,unit in groups])
for ws in (idx,full):style(ws);ws.freeze_panes='I5' if ws is full else 'A5';ws.row_dimensions[1].height=25;ws.row_dimensions[2].height=30
for c,width in {'A':18,'B':18,'C':16,'D':48,'E':16,'F':28,'G':42,'H':56,'I':25}.items():idx.column_dimensions[c].width=width
for c,width in enumerate([18,18,16,48,16,28,42,28]+[18]*len(groups),1):full.column_dimensions[get_column_letter(c)].width=width
full.row_dimensions[4].height=70
lab.save(OUT/'03_Лаборатория_20_скважин.xlsx')
package=OUT/'ROSA_Учебный_набор_20_скважин_797-RC.zip'
with zipfile.ZipFile(package,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=7) as archive:
 for name in ('01_Техрежим_20_скважин.xlsx','02_Реестр_отказов_20_скважин.xlsx','03_Лаборатория_20_скважин.xlsx','04_Экономика_20_скважин.xlsx'):
  archive.write(OUT/name,name)
 for name in ('РОСА_БЛАНК_ЛАБОРАТОРНОГО_АНАЛИЗА.xlsx','РОСА_БЛАНК_РЕЕСТРА_ОТКАЗОВ_632-RC.xlsx','РОСА_БЛАНК_ЭКОНОМИКИ.xlsx'):
  archive.write(ROOT/name,name)
print(json.dumps({'wells':len(wells),'failures':sum(len(w['failures']) for w in wells),'labPasses':sum(len(w['labPasses']) for w in wells),'labCatalogFields':len(groups),'economicsScenarios':60*3,'package':str(package),'out':str(OUT)},ensure_ascii=False))
